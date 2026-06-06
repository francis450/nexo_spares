"""
Data migration script for Nexo Auto Spares.
Imports items from two Excel files, creates Item Prices and opening Stock Entries.

Run with:
    bench --site nexoautospares.erpkenya.com execute \
        nexo_spares.install.import_items.run_import

Idempotent — skips items that already exist by item_code.
"""

import os
import re

import frappe
import openpyxl
from frappe.utils import flt, nowdate

BENCH_PATH = "/home/frappeuser/frappe-bench"
SPOILERS_FILE = os.path.join(BENCH_PATH, "Nexo spoilers.xlsx")
SECOND_STOCK_FILE = os.path.join(BENCH_PATH, "NEXO AUTO second stock.xlsx")

DEFAULT_WAREHOUSE = "Stores - NA"
DEFAULT_COMPANY = "Nexo Autospares"
PRICE_LIST = "Standard Selling"

# (name, parent_item_group) — map to existing site groups where sensible
ITEM_GROUPS = [
    ("Auto Accessories", "All Item Groups"),
    ("Chrome Kits",      "Car Accessories"),
    ("Floor Mats",       "Car Accessories"),
    ("Fog Lamps",        "Lighting"),
    ("Horns",            "Car Accessories"),
    ("LED Lighting",     "Lighting"),
    ("Spoilers",         "Car Accessories"),
    ("Windbreakers",     "Car Accessories"),
    ("Wing Mirrors",     "Car Accessories"),
]


def make_item_code(description):
    code = re.sub(r"[^A-Za-z0-9 \-/]", "", description.strip())
    code = re.sub(r"\s+", "-", code.strip()).upper()
    return code[:50]


def infer_item_group(description):
    d = description.upper()
    if "WINDBREAKER" in d:
        return "Windbreakers"
    if "FOG LAMP" in d or "FOG-LAMP" in d or "FOG  LAMP" in d:
        return "Fog Lamps"
    if "LED" in d or "COB" in d or "BULB" in d:
        return "LED Lighting"
    if "FLOOR MAT" in d or "DASHBOARD" in d:
        return "Floor Mats"
    if "CHROM" in d or "GRILL" in d:
        return "Chrome Kits"
    if "HORN" in d:
        return "Horns"
    if "WING MIRROR" in d:
        return "Wing Mirrors"
    if "SPOILER" in d:
        return "Spoilers"
    return "Auto Accessories"


def ensure_item_groups():
    for group_name, parent in ITEM_GROUPS:
        if not frappe.db.exists("Item Group", group_name):
            doc = frappe.get_doc({
                "doctype": "Item Group",
                "item_group_name": group_name,
                "parent_item_group": parent,
            })
            doc.insert(ignore_permissions=True)
            frappe.db.commit()
            print(f"  Created Item Group: {group_name} (under {parent})")


def create_item(item_code, item_name, item_group, valuation_rate, selling_price, qty, uom="Nos"):
    if frappe.db.exists("Item", item_code):
        print(f"  SKIP (exists): {item_code}")
        return False

    print(f"  CREATE: {item_code} | {item_group} | cost={flt(valuation_rate,0)} | sell={flt(selling_price,0)} | qty={qty}")

    item = frappe.get_doc({
        "doctype": "Item",
        "item_code": item_code,
        "item_name": item_name,
        "item_group": item_group,
        "stock_uom": uom,
        "is_stock_item": 1,
        "valuation_method": "FIFO",
        "valuation_rate": flt(valuation_rate, 2),
        # Do NOT set standard_rate here — ERPNext auto-creates an Item Price from it,
        # causing a duplicate when we insert our own Item Price below.
        "item_defaults": [{
            "company": DEFAULT_COMPANY,
            "default_warehouse": DEFAULT_WAREHOUSE,
        }],
    })
    item.insert(ignore_permissions=True)

    if flt(selling_price) > 0:
        ip = frappe.get_doc({
            "doctype": "Item Price",
            "item_code": item_code,
            "price_list": PRICE_LIST,
            "selling": 1,
            "buying": 0,
            "price_list_rate": flt(selling_price, 2),
            "currency": "KES",
        })
        ip.insert(ignore_permissions=True)

    if flt(qty) > 0:
        se = frappe.get_doc({
            "doctype": "Stock Entry",
            "stock_entry_type": "Material Receipt",
            "is_opening": "Yes",
            "company": DEFAULT_COMPANY,
            "posting_date": nowdate(),
            "items": [{
                "item_code": item_code,
                "qty": flt(qty),
                "basic_rate": flt(valuation_rate, 2),
                "t_warehouse": DEFAULT_WAREHOUSE,
                # Must be Balance Sheet (Asset/Liability) when is_opening=Yes
                "expense_account": "Temporary Opening - NA",
            }],
        })
        se.insert(ignore_permissions=True)
        se.submit()

    frappe.db.commit()
    return True


def parse_spoilers():
    """Parse Nexo spoilers.xlsx — 6 spoiler items."""
    print("\n=== Nexo spoilers.xlsx ===")
    wb = openpyxl.load_workbook(SPOILERS_FILE, data_only=True)
    ws = wb.active
    items = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 3 or i > 8:
            continue
        desc = row[0]
        qty = row[1]
        valuation = row[8]   # col I: laded cost per item
        selling = row[9]     # col J: selling price
        if not desc or str(desc).strip().upper() in ("TOTAL", ""):
            continue
        name = f"Spoiler {str(desc).strip()}"
        items.append({
            "description": name,
            "qty": qty,
            "valuation_rate": valuation,
            "selling_price": selling,
            "uom": "Nos",
            "item_group": "Spoilers",
        })
    print(f"  Found {len(items)} items")
    return items


def parse_second_stock():
    """Parse NEXO AUTO second stock.xlsx — sheet 03-06, ~58 items."""
    print("\n=== NEXO AUTO second stock.xlsx ===")
    wb = openpyxl.load_workbook(SECOND_STOCK_FILE, data_only=True)

    # Try to get the correct sheet
    sheet_name = "03-06"
    if sheet_name not in wb.sheetnames:
        sheet_name = wb.sheetnames[0]
        print(f"  Warning: sheet '03-06' not found, using '{sheet_name}'")

    ws = wb[sheet_name]
    items = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i < 6:
            continue
        item_num = row[0]
        desc = row[1]
        qty = row[2]
        uom_raw = row[3]
        valuation = row[9]   # col J: laded cost/item
        selling = row[10]    # col K: selling price

        if not desc or not item_num:
            continue
        desc_str = str(desc).strip()
        if not desc_str or desc_str.upper().startswith("TOTAL"):
            continue

        uom_map = {"SET": "Set", "PC": "Nos", "PCS": "Nos", "ROLL": "Roll", "PAIR": "Pair", "NOS": "Nos", "UNIT": "Nos"}
        uom = uom_map.get(str(uom_raw).strip().upper(), "Nos") if uom_raw else "Nos"

        items.append({
            "description": desc_str,
            "qty": qty,
            "valuation_rate": valuation,
            "selling_price": selling,
            "uom": uom,
            "item_group": infer_item_group(desc_str),
        })

    print(f"  Found {len(items)} items")
    return items


@frappe.whitelist()
def run_import():
    """
    Main entry point.
    bench --site nexoautospares.erpkenya.com execute nexo_spares.install.import_items.run_import
    """
    print("\n" + "=" * 50)
    print("Nexo Spares — Data Import")
    print("=" * 50)
    print(f"Company  : {DEFAULT_COMPANY}")
    print(f"Warehouse: {DEFAULT_WAREHOUSE}")
    print(f"PriceList: {PRICE_LIST}")

    if not frappe.db.exists("Warehouse", DEFAULT_WAREHOUSE):
        frappe.throw(
            f"Warehouse '{DEFAULT_WAREHOUSE}' not found. "
            "Verify the exact name in ERPNext and update DEFAULT_WAREHOUSE in the script."
        )
    if not frappe.db.exists("Company", DEFAULT_COMPANY):
        frappe.throw(
            f"Company '{DEFAULT_COMPANY}' not found. "
            "Verify the exact name and update DEFAULT_COMPANY in the script."
        )
    if not frappe.db.exists("Price List", PRICE_LIST):
        frappe.throw(
            f"Price List '{PRICE_LIST}' not found. "
            "Verify the exact name and update PRICE_LIST in the script."
        )

    ensure_item_groups()

    all_items = parse_spoilers() + parse_second_stock()
    print(f"\nTotal items to process: {len(all_items)}")

    created = 0
    skipped = 0
    for item_data in all_items:
        item_code = make_item_code(item_data["description"])
        result = create_item(
            item_code=item_code,
            item_name=item_data["description"],
            item_group=item_data["item_group"],
            valuation_rate=item_data["valuation_rate"],
            selling_price=item_data["selling_price"],
            qty=item_data["qty"],
            uom=item_data["uom"],
        )
        if result:
            created += 1
        else:
            skipped += 1

    print(f"\n{'=' * 50}")
    print(f"Import complete: {created} created, {skipped} skipped")
    print("=" * 50)

    return {"status": "ok", "created": created, "skipped": skipped}
